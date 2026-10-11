import zipfile, hashlib, os, sys

DOCX = r"C:\laragon\www\majapahit\seruti\3509_Laporan_Penyelesaian_Seruti_TW3_2026.docx"
BASE = r"C:\laragon\www\majapahit\seruti"

SRC = {
    "85d91c7427ac": "LOGO BPS (resmi)",
    "41448d5493bc": "TANDA TANGAN (scan)",
    "c8c0a237b801": "DASHBOARD INTI",
    "f91561cc0b0e": "DASHBOARD MAK",
}

print("=" * 68)
print("VERIFIKASI DOKUMEN - 3509_Laporan_Penyelesaian_Seruti_TW3_2026.docx")
print("=" * 68)

z = zipfile.ZipFile(DOCX)
print("ZIP integrity :", "OK (tidak ada member korup)" if z.testzip() is None else "CORRUPT")

imgs = sorted(n for n in z.namelist() if n.startswith("word/media/"))
print("Media embed   :", len(imgs), "gambar")
for n in imgs:
    d = z.read(n)
    h = hashlib.sha1(d).hexdigest()[:12]
    print("   {:20s} {:8d} B  {}  {}".format(n, len(d), h, SRC.get(h, "TIDAK DIKENAL")))

unknown = [n for n in imgs if hashlib.sha1(z.read(n)).hexdigest()[:12] not in SRC]
print("Gambar asing   :", "TIDAK ADA" if not unknown else unknown)

# --- struktur dokumen ---
from docx import Document
d = Document(DOCX)
print()
print("Paragraf      :", len(d.paragraphs))
print("Tabel         :", len(d.tables))
print("Inline shapes :", len(d.inline_shapes))
print("Page break    :", "ada" if any(
    "w:br" in r._element.xml and 'type="page"' in r._element.xml
    for p in d.paragraphs for r in p.runs) else "tidak terdeteksi")

print()
print("--- ISI TABEL PENGESAHAN ---")
for t in d.tables:
    txt = t.rows[0].cells[-1].text
    if "NANANG" in txt:
        for line in txt.split("\n"):
            if line.strip():
                print("   ", line)
        print("    [gambar TTD di dalam sel :",
              "YA" if "<w:drawing" in t.rows[0].cells[-1]._tc.xml else "TIDAK", "]")

print()
print("--- CATATAN TABEL 5 ---")
for p in d.paragraphs:
    if p.text.strip().startswith("Catatan:"):
        print("   ", p.text)

print()
print("--- HEADING ---")
for p in d.paragraphs:
    t = p.text.strip()
    if t and t[0].isalpha() and (t.isupper() or t[:2] in ("4.", "I.", "V.")) and len(t) < 90:
        r = p.runs[0] if p.runs else None
        sz = r.font.size.pt if (r and r.font.size) else "?"
        if sz and sz >= 10:
            print("   {:>5} pt  {}".format(str(sz), t))
