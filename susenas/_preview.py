import fitz
from docx import Document

doc = Document(r'C:/laragon/www/majapahit/susenas/3509_Laporan_Pengolahan_Susenas_September_2026_50Persen.docx')

print('--- bullet / special char check ---')
for i in (35, 36, 37):
    print(i, repr(doc.paragraphs[i].text[:60]))

print('--- cover lines repr ---')
for i in (14, 15, 16, 17):
    print(i, repr(doc.paragraphs[i].text))

pdf = r'C:/laragon/www/majapahit/susenas/3509_Laporan_Pengolahan_Susenas_September_2026_50Persen.pdf'
d = fitz.open(pdf)
print('\nPDF pages:', d.page_count)
out = r'C:/laragon/www/majapahit/susenas/_preview'
import os, shutil
shutil.rmtree(out, ignore_errors=True)
os.makedirs(out, exist_ok=True)
blank = []
for i in range(d.page_count):
    page = d[i]
    pix = page.get_pixmap(dpi=110)
    pix.save(f'{out}/page_{i+1:02d}.png')
    r = page.rect
    blocks = [b for b in page.get_text('blocks') if b[4].strip()]
    over = [b for b in blocks if b[2] > r.x1 - 4 or b[0] < 4]
    if len(blocks) <= 1:
        blank.append(i + 1)
    first = (blocks[0][4].strip().replace('\n', ' ')[:60]) if blocks else '(kosong)'
    print(f'  p{i+1:02d}: blocks={len(blocks):2d} overflow={len(over)} | {first}')
print('BLANK PAGES:', blank if blank else 'tidak ada')
print('rendered ->', out)
