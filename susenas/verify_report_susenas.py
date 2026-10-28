# -*- coding: utf-8 -*-
"""
Verifikasi laporan Susenas September 2026 - BPS Kabupaten Jember (3509).

Memastikan seluruh teks, tabel, gambar, dan metadata dokumen terisi presisi,
konsisten dengan sumber data monitoring, dan bebas dari sisa template lama.

Dipakai  : C:\\laragon\\bin\\python\\python-3.10\\python.exe
Jalankan : python verify_report_susenas.py
"""

from __future__ import annotations

import re
import sys
from pathlib import Path

from docx import Document

BASE = Path(r"C:\laragon\www\majapahit\susenas")
DATA_DIR = BASE / "data-sm2-2026"
DOCX = BASE / "3509_Laporan_Pengolahan_Susenas_September_2026_50Persen.docx"
PDF = BASE / "3509_Laporan_Pengolahan_Susenas_September_2026_50Persen.pdf"
XLSX = DATA_DIR / "rekap_progres_pengolahan_3509.xlsx"
TEMPLATE = BASE / "template" / "3509_Laporan_Pengolahan_Susenas_September_2026-rev.docx"

FILE_K = DATA_DIR / "25-09-2026_154754-Progress_Pengolahan_DokM_3509.xls"
FILE_KP = DATA_DIR / "25-09-2026_154833-Progress_Pengolahan_DokKP_3509.xls"
FILE_LK = DATA_DIR / "LK Pengolahan Sampel.xlsx"

results = []


def check(name, ok, detail=""):
    results.append((bool(ok), name, detail))
    print(f"  [{'PASS' if ok else 'FAIL'}] {name}" + (f" - {detail}" if detail else ""))
    return bool(ok)


def parse_source(path: Path):
    """Ekstrak tabel progres dari berkas ekspor Web Monitoring (HTML)."""
    html = path.read_text(encoding="utf-8")
    cells = re.findall(r"<t[dh][^>]*>(.*?)</t[dh]>", html, re.S)
    vals = [re.sub(r"\s+", " ", re.sub(r"<[^>]+>", " ", c)).strip() for c in cells]
    rows, i = [], 0
    while i < len(vals):
        if vals[i].startswith("TOTAL"):
            break
        if vals[i].isdigit() and i + 6 < len(vals):
            kode = re.search(r"\((\d+)\)", vals[i + 1])
            rows.append({
                "no": int(vals[i]),
                "kode": kode.group(1) if kode else vals[i + 1],
                "target": int(vals[i + 2]),
                "clean": int(vals[i + 3]),
                "clean_pct": vals[i + 4],
                "error": int(vals[i + 5]),
                "error_pct": vals[i + 6],
            })
            i += 7
        else:
            i += 1
    return rows


def table_text(table):
    return "\n".join("\t".join(c.text for c in row.cells) for row in table.rows)


def parse_lk():
    """Agregat independen dari LK Pengolahan Sampel untuk verifikasi silang."""
    import openpyxl
    wb = openpyxl.load_workbook(FILE_LK, data_only=True)
    lk = {}
    rows = [r for r in wb["Rekap_Kirim_Dokumen"].iter_rows(min_row=4, values_only=True)
            if r[2] is not None]
    lk["kirim_total"] = len(rows)
    lk["kirim_sudah"] = sum(1 for r in rows if str(r[4]).lower() == "sudah")
    lk["kirim_belum"] = lk["kirim_total"] - lk["kirim_sudah"]
    rek = [r for r in wb["Rekap"].iter_rows(min_row=2, values_only=True) if r[2]]
    lk["entry_kor"] = sum(1 for r in rek if str(r[15]) in ("1", "1.0"))
    lk["entry_kp"] = sum(1 for r in rek if str(r[14]) in ("1", "1.0"))
    from collections import Counter
    lk["beban"] = dict(Counter(str(r[9]).strip() for r in rek if r[9]).most_common())
    sheets = ["Nur Ida Suryandari", "Iffa Dzakiyya Khairunnisa", "Laviana Ika Putrisari",
              "Anung Anindhita Pratiwi", "Aminatus Sholeha", "Putri Salsabhila Fahira",
              "Astri Widarianti", "Prasistiwi Andrianingtyas"]
    lk["rt_catatan"] = sum(
        1 for nm in sheets
        for r in wb[nm].iter_rows(min_row=2, values_only=True)
        if r[1] and str(r[1]).strip()
        and any(r[i] for i in (4, 5, 6, 7, 8, 9)))
    return lk


def main():
    print("=" * 74)
    print("VERIFIKASI LAPORAN SUSENAS SEPTEMBER 2026 - BPS KABUPATEN JEMBER (3509)")
    print("=" * 74)

    # ---------- 0. KEBERADAAN BERKAS ----------
    print("\n0. BERKAS KELUARAN")
    check("Template adaptif tersedia", TEMPLATE.exists(), TEMPLATE.name)
    check("Dokumen DOCX dibuat", DOCX.exists(),
          f"{DOCX.stat().st_size:,} bytes" if DOCX.exists() else "tidak ada")
    check("Dokumen PDF dibuat", PDF.exists(),
          f"{PDF.stat().st_size:,} bytes" if PDF.exists() else "tidak ada")
    check("Rekap XLSX (jejak audit) dibuat", XLSX.exists())
    if not DOCX.exists():
        print("\nDokumen tidak ditemukan - verifikasi dihentikan.")
        return 1

    doc = Document(str(DOCX))
    paras = [p.text for p in doc.paragraphs]
    full = "\n".join(paras)
    all_text = full + "\n" + "\n".join(table_text(t) for t in doc.tables)
    tables_txt = "\n".join(table_text(t) for t in doc.tables)

    # ---------- sumber data ----------
    src_k = parse_source(FILE_K)
    src_kp = parse_source(FILE_KP)
    n_bs = len(src_k)
    tgt_k = sum(r["target"] for r in src_k)
    tgt_kp = sum(r["target"] for r in src_kp)
    cl_k = sum(r["clean"] for r in src_k)
    er_k = sum(r["error"] for r in src_k)
    cl_kp = sum(r["clean"] for r in src_kp)
    er_kp = sum(r["error"] for r in src_kp)
    blm_k = tgt_k - cl_k - er_k
    blm_kp = tgt_kp - cl_kp - er_kp

    print("\n1. KONSISTENSI DATA SUMBER vs DOKUMEN")
    check(f"Jumlah blok sensus source = {n_bs}", n_bs == 28)
    check(f"Total target source = {tgt_k} / {tgt_kp} RT",
          tgt_k == 280 and tgt_kp == 280)
    check(f"Total clean Dok K = {cl_k}", cl_k == 8)
    check(f"Total clean Dok KP = {cl_kp}", cl_kp == 40)
    check(f"Total error Dok K = {er_k}", er_k == 261)
    check(f"Total error Dok KP = {er_kp}", er_kp == 227)
    check(f"Belum diproses Dok K = {blm_k}", blm_k == 11)
    check(f"Belum diproses Dok KP = {blm_kp}", blm_kp == 13)

    # ---------- 2. HALAMAN JUDUL / KOP ----------
    print("\n2. HALAMAN JUDUL / KOP")
    cover = "\n".join(paras[:24])
    for text in [
        "BADAN PUSAT STATISTIK",
        "KABUPATEN JEMBER",
        "PROVINSI JAWA TIMUR",
        "LAPORAN KEGIATAN",
        "PENYELESAIAN PENGOLAHAN SAMPEL SUSENAS",
        "SEPTEMBER 2026 (PROGRESS 50%)",
        "Periode Pelaporan",
        "September 2026",
        "Kantor BPS Kabupaten Jember (3509)",
        "28 Blok Sensus (280 Sampel Rumah Tangga)",
        "50% (140 Sampel RT / 14 Blok Sensus Clean)",
        "Tim Pengolahan dan Layanan Statistik",
        "BPS Kabupaten Jember",
        "JEMBER, 2026",
    ]:
        check(f"Cover memuat: {text!r}", text in cover)

    # ---------- 3. STRUKTUR BAB ----------
    print("\n3. STRUKTUR BAB (URUTAN BAKU BPS)")
    expected = [
        "I. PENDAHULUAN",
        "II. INFORMASI UMUM",
        "III. MAKSUD DAN TUJUAN",
        "IV. DASAR HUKUM",
        "V. RUANG LINGKUP DAN METODE PENGOLAHAN",
        "VI. REKAPITULASI BEBAN WILAYAH DAN SAMPEL",
        "VII. HASIL DAN CAPAIAN PENGOLAHAN 50%",
        "VIII. KENDALA, SOLUSI, KESIMPULAN, DAN REKOMENDASI",
        "LEMBAR PENGESAHAN",
        "LAMPIRAN",
    ]
    positions = []
    for ch in expected:
        idx = next((i for i, p in enumerate(paras) if p.strip() == ch), None)
        positions.append(idx)
        check(f"Bab ada: {ch}", idx is not None)
    known = [p for p in positions if p is not None]
    check("Seluruh bab berurutan sesuai urutan baku", known == sorted(known),
          f"posisi={positions}")

    for sub in ["A. Maksud", "B. Tujuan", "A. Ruang Lingkup",
                "B. Metode dan Tahapan Pengolahan",
                "C. Waktu dan Tempat Pelaksanaan",
                "A. Kendala dan Solusi", "B. Kesimpulan", "C. Rekomendasi"]:
        check(f"Sub-bab ada: {sub}", any(p.strip() == sub for p in paras))

    # ---------- 4. LEMBAR PENGESAHAN ----------
    print("\n4. LEMBAR PENGESAHAN")
    check("Tanggal pengesahan", "Jember, 30 September 2026" in tables_txt)
    check("Jabatan penandatangan",
          "Ketua Tim Pengolahan dan Layanan Statistik" in tables_txt)
    check("Nama penandatangan", "Nanang Pamungkas" in tables_txt)
    check("NIP penandatangan", "198505172011011018" in tables_txt)

    # ---------- 5. DASAR HUKUM ----------
    print("\n5. DASAR HUKUM")
    for d in ["Undang-Undang Nomor 16 Tahun 1997",
              "Peraturan Pemerintah Nomor 51 Tahun 1999",
              "Peraturan Badan Pusat Statistik",
              "Surat Keputusan Kepala BPS Kabupaten Jember"]:
        check(f"Memuat: {d}", d in all_text)

    # ---------- 6. STRUKTUR TABEL ----------
    print("\n6. STRUKTUR TABEL")
    check(f"Jumlah tabel = 9 (ditemukan {len(doc.tables)})",
          len(doc.tables) == 9)
    if len(doc.tables) != 9:
        for i, t in enumerate(doc.tables):
            print(f"      T{i}: {len(t.rows)}x{len(t.columns)} | "
                  f"{t.rows[0].cells[0].text!r}")
        print("\nRingkasan verifikasi dihentikan (struktur tabel tidak sesuai).")
        failed = sum(1 for ok, _, _ in results if not ok)
        return 1 if failed else 0

    t_info, t_bab6, t_bab7a, t_bab7b, t_kend, _t_ttd, t_l1, t_l2, t_l4 = doc.tables

    check(f"Tabel Informasi Umum 19x2 (ditemukan "
          f"{len(t_info.rows)}x{len(t_info.columns)})",
          len(t_info.rows) == 19 and len(t_info.columns) == 2)
    info = table_text(t_info)
    for v in ["[3509] Kabupaten Jember", "28 Blok Sensus (BS)",
              "280 Sampel Rumah Tangga (RT)", "50% (140 RT / 14 BS Clean)",
              "8 RT clean (2,86%)", "40 RT clean (14,29%)"]:
        check(f"  Info umum memuat: {v}", v in info)

    # ---- Tabel Bab VI : rekap 28 blok sensus ----
    check(f"Tabel Bab VI = 30 baris (ditemukan {len(t_bab6.rows)})",
          len(t_bab6.rows) == 30)
    check(f"Tabel Bab VI = 8 kolom (ditemukan {len(t_bab6.columns)})",
          len(t_bab6.columns) == 8)
    mism = [s["kode"] for i, s in enumerate(src_k)
            if t_bab6.rows[i + 1].cells[1].text.strip() != s["kode"]
            or t_bab6.rows[i + 1].cells[3].text.strip() != str(s["clean"])
            or t_bab6.rows[i + 1].cells[4].text.strip() != str(s["error"])]
    check("  28 baris Bab VI cocok dengan sumber Dok K", not mism,
          f"{len(mism)} selisih")
    mism2 = [s["kode"] for i, s in enumerate(src_kp)
             if t_bab6.rows[i + 1].cells[5].text.strip() != str(s["clean"])
             or t_bab6.rows[i + 1].cells[6].text.strip() != str(s["error"])]
    check("  28 baris Bab VI cocok dengan sumber Dok KP", not mism2,
          f"{len(mism2)} selisih")
    last = t_bab6.rows[29]
    check("  Baris TOTAL Bab VI benar",
          last.cells[1].text.strip() == "TOTAL"
          and last.cells[2].text.strip() == "280"
          and last.cells[3].text.strip() == str(cl_k)
          and last.cells[5].text.strip() == str(cl_kp))

    # ---- Tabel Bab VII-a : indikator milestone ----
    check(f"Tabel indikator milestone = 15 baris (ditemukan {len(t_bab7a.rows)})",
          len(t_bab7a.rows) == 15)
    info7 = table_text(t_bab7a)
    for v in ["50% = 140 RT clean / 14 BS clean",
              "8 RT dari 280 RT (2,86%)",
              "40 RT dari 280 RT (14,29%)",
              "48 RT dari 560 record (8,57%)",
              "132 RT lagi", "100 RT lagi", "BELUM TERCAPAI"]:
        check(f"  Milestone memuat: {v}", v in info7)

    # ---- Tabel Bab VII-b : clean vs error ----
    check(f"Tabel clean vs error = 4 baris (ditemukan {len(t_bab7b.rows)})",
          len(t_bab7b.rows) == 4)
    txt7 = table_text(t_bab7b)
    for v in ["2,86%", "93,21%", "3,93%", "14,29%", "81,07%", "4,64%",
              "8,57%", "87,14%", "4,29%"]:
        check(f"  Persentase memuat: {v}", v in txt7)

    # ---- Tabel kendala/solusi ----
    check(f"Tabel kendala/solusi = 7 baris (ditemukan {len(t_kend.rows)})",
          len(t_kend.rows) == 7)

    # ---- Lampiran 1 & 2 ----
    for nm, t, src, cl, er in [("Lampiran 1 (Dok K)", t_l1, src_k, cl_k, er_k),
                               ("Lampiran 2 (Dok KP)", t_l2, src_kp, cl_kp, er_kp)]:
        check(f"{nm} = 30x8 (ditemukan {len(t.rows)}x{len(t.columns)})",
              len(t.rows) == 30 and len(t.columns) == 8)
        mm = [s["kode"] for i, s in enumerate(src)
              if t.rows[i + 1].cells[1].text.strip() != s["kode"]
              or t.rows[i + 1].cells[3].text.strip() != str(s["clean"])
              or t.rows[i + 1].cells[5].text.strip() != str(s["error"])
              or t.rows[i + 1].cells[4].text.strip() != s["clean_pct"]]
        check(f"  {nm}: 28 baris cocok sumber", not mm, f"{len(mm)} selisih")
        lt = t.rows[29]
        check(f"  {nm}: baris TOTAL benar",
              lt.cells[3].text.strip() == str(cl)
              and lt.cells[5].text.strip() == str(er))

    # ---- Lampiran 4 : alokasi beban per pengolah ----
    check(f"Lampiran 4 = 10 baris (ditemukan {len(t_l4.rows)})",
          len(t_l4.rows) == 10)
    check(f"Lampiran 4 = 4 kolom (ditemukan {len(t_l4.columns)})",
          len(t_l4.columns) == 4)
    l4_txt = table_text(t_l4)
    for nm_p in ["Anung Anindhita P", "Prasistiwi", "Iffa Dzakiyya Khairunnisa",
                 "Laviana Ika Putrisari", "Aminatus Sholeha", "Putri Salsabhila Fahira",
                 "Nur Ida Suryandari", "Astri Widarianti"]:
        check(f"  Lampiran 4 memuat pengolah: {nm_p}", nm_p in l4_txt)
    last_l4 = t_l4.rows[9]
    check("  Lampiran 4 baris TOTAL = 280",
          last_l4.cells[1].text.strip() == "TOTAL"
          and last_l4.cells[2].text.strip() == "280"
          and last_l4.cells[3].text.strip() == "100,00%")

    # ---------- 6b. KONSISTENSI DATA LEMBAR KERJA (LK) ----------
    print("\n6b. KONSISTENSI DATA LEMBAR KERJA (LK) PENGOLAHAN")
    lk = parse_lk()
    check("LK: Total kirim = 280 RT", lk["kirim_total"] == 280)
    check("LK: Sudah kirim = 242 RT", lk["kirim_sudah"] == 242)
    check("LK: Belum kirim = 38 RT", lk["kirim_belum"] == 38)
    check("LK: Entry Dokumen Kor = 155 RT", lk["entry_kor"] == 155)
    check("LK: Entry Dokumen KP = 140 RT", lk["entry_kp"] == 140)
    check("LK: Jumlah RT catatan editing = 67 RT", lk["rt_catatan"] == 67)
    check("LK: 8 pengolah sampel terdaftar", len(lk["beban"]) == 8)
    for fig in ["242 dari 280 RT", "86,43%", "38 RT belum",
                "155 RT untuk Dokumen K", "140 RT untuk Dokumen KP",
                "67 RT", "Isian wajib kosong/terlewat (43 RT)",
                "Koordinat tidak sesuai (29 RT)",
                "Isian di luar range (18 RT)",
                "Selisih hasil hitung form vs aplikasi (10 RT)"]:
        check(f"  Naratif LK termuat: {fig}", fig in all_text)

    # ---------- 7. BUKTI VISUAL ----------
    print("\n7. BUKTI VISUAL MONITORING")
    check(f"Jumlah gambar = 3 (ditemukan {len(doc.inline_shapes)})",
          len(doc.inline_shapes) == 3)
    for cap in ["Gambar 1. Monitoring Pengolahan Dokumen K",
                "Gambar 2. Monitoring Pengolahan Dokumen KP"]:
        check(f"Keterangan: {cap}", cap in all_text)

    # ---------- 8. ISI NARATIF KUNCI ----------
    print("\n8. ISI NARATIF KUNCI")
    for text in [
        "Survei Konsumsi Rumah Tangga (Susenas) Semester 2 Tahun 2026",
        "Seruti Triwulan III Tahun 2026",
        "Kode Wilayah 3509",
        "webmonitoring.bps.go.id",
        "Penerimaan Dokumen",
        "Editing dan Coding",
        "Entri Data ke Web Monitoring Susenas",
        "Rule Validation",
        "Resolusi Error Dokumen K dan Dokumen KP",
        "Evaluasi Anomali Isian Konsumsi dan Kor",
        "Lampiran 3. Tangkapan Layar Monitoring Web Susenas",
    ]:
        check(f"Memuat: {text}", text in all_text)

    # ---------- 9. BEBAS SISA TEMPLATE LAMA ----------
    print("\n9. BEBAS SISA TEMPLATE LAMA (SAKERNAS / SE2026)")
    for text in ["ASSIGNMENT WILAYAH PPL DAN PML", "PERSIAPAN SENSUS EKONOMI 2026",
                 "Sensus Ekonomi 2026", "SE2026", "1.156.276", "31 Kecamatan",
                 "Jumlah PML", "Status Assignment", "have-not-assigned",
                 "April s.d. Juni 2026", "Juni 2026"]:
        check(f"Tidak memuat sisa: {text!r}", text not in all_text)

    # ---------- 10. PDF ----------
    print("\n10. PDF")
    if PDF.exists():
        raw = PDF.read_bytes()
        check("Header PDF valid", raw[:5] == b"%PDF-", repr(raw[:5]))
        check("PDF memuat objek halaman",
              b"/Type/Page" in raw or b"/Type /Page" in raw)
        check("PDF tidak kosong", PDF.stat().st_size > 50_000,
              f"{PDF.stat().st_size:,} bytes")
    else:
        check("PDF tersedia", False)

    # ---------- RINGKASAN ----------
    passed = sum(1 for ok, _, _ in results if ok)
    failed = sum(1 for ok, _, _ in results if not ok)
    print("\n" + "=" * 74)
    print(f"RINGKASAN VERIFIKASI : {passed} LULUS / {failed} GAGAL "
          f"dari {len(results)} pemeriksaan")
    print("=" * 74)
    if failed:
        print("\nDaftar pemeriksaan yang GAGAL:")
        for ok, name, detail in results:
            if not ok:
                print(f"  - {name} {detail}")
    return 1 if failed else 0


if __name__ == "__main__":
    sys.exit(main())
