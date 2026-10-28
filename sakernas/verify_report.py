from docx import Document

doc = Document('C:/laragon/www/majapahit/sakernas/3509_Laporan_Assignment_PPL_PML_Sakernas_2026.docx')

print(f"Total paragraphs: {len(doc.paragraphs)}")
print(f"Total tables: {len(doc.tables)}")

print("\n=== ALL PARAGRAPHS ===")
for i, para in enumerate(doc.paragraphs):
    text = para.text
    if text.strip():
        print(f"P{i:03d}: {text[:250]}")

print("\n=== TABLE 0 - INFORMASI UMUM ===")
table0 = doc.tables[0]
for r_idx, row in enumerate(table0.rows):
    cells = [cell.text for cell in row.cells]
    print(f"  Row {r_idx}: {cells}")

print("\n=== TABLE 1 - DAFTAR WILAYAH ===")
table1 = doc.tables[1]
for r_idx, row in enumerate(table1.rows):
    cells = [cell.text for cell in row.cells]
    print(f"  Row {r_idx}: {cells}")

print("\n=== TABLE 2 - HASIL CAPAIAN ===")
table2 = doc.tables[2]
for r_idx, row in enumerate(table2.rows):
    cells = [cell.text for cell in row.cells]
    print(f"  Row {r_idx}: {cells}")

print("\n=== TABLE 3 - TANDA TANGAN ===")
table3 = doc.tables[3]
for r_idx, row in enumerate(table3.rows):
    cells = [cell.text for cell in row.cells]
    print(f"  Row {r_idx}: {repr(cells)}")

print("\n=== TABLE 4 - LAMPIRAN ===")
table4 = doc.tables[4]
for r_idx, row in enumerate(table4.rows):
    cells = [cell.text for cell in row.cells]
    print(f"  Row {r_idx}: {cells}")
